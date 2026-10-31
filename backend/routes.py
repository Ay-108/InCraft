"""
InCraft Routes
--------------
Application routes and request handling for:

- Authentication
- Dashboard
- Inventory
- Products
- Sales
- Analytics
- Activity / History
- Settings
- Search
- Import / Export
"""

from datetime import datetime, timedelta
from uuid import UUID
from functools import wraps
from typing import Any, Callable, ParamSpec, cast

from flask import (
    Blueprint,
    abort,
    render_template,
    request,
    redirect,
    url_for,
    jsonify,
    flash,
    session
)
from flask.typing import ResponseReturnValue

from werkzeug.security import (
    generate_password_hash,
    check_password_hash
)

from backend.db import (
    fetch_one,
    fetch_all,
    execute_query,
    execute_many,
    execute_transaction,
    get_database_status
)

from backend.config import (
    DATABASE_CONFIGURED,
    MISSING_DATABASE_ENV_VARS,
    SECRET_KEY_CONFIGURED,
    IS_PRODUCTION,
)

from backend.ml.forecast import forecast_product
from backend.ml.stockout import assess_stockout
from backend.ml.reorder import recommend_reorder

from backend.utils import (
    allowed_file,
    generate_csv_response,
    validate_email,
    validate_password,
    validate_required_fields,
    validate_positive_number,
    validate_non_negative_number,
    get_stock_status,
    calculate_stock_value,
    calculate_sale_total,
    format_currency,
    login_user,
    logout_user,
    login_required,
    get_current_user_id,
    create_activity_description,
    normalize_search_query,
    generate_reference,
    safe_int,
    safe_float
)


# ============================================================
# BLUEPRINT
# ============================================================

main_bp = Blueprint(
    "main",
    __name__
)

P = ParamSpec("P")


# ============================================================
# ROLE-BASED ACCESS CONTROL
# ============================================================

def get_current_user_role():
    """Return the signed-in user's current role from the database.

    Never trust a role submitted by the browser or stored only in a session.
    Authorization checks fail closed if the role cannot be verified.
    """
    user_id = get_current_user_id()
    if not user_id:
        return None

    try:
        current_user = fetch_one(
            """
            SELECT role, status
            FROM users
            WHERE user_id = %s
            LIMIT 1
            """,
            (user_id,)
        )
    except Exception as error:
        print(f"User-role lookup failed: {error}")
        abort(503)

    if current_user is None:
        logout_user()
        return None

    # Inactive accounts must lose access even if an older session still exists.
    account_status = str(cast(Any, current_user).get("status") or "").strip().lower()
    if account_status != "active":
        logout_user()
        return None

    return str(cast(Any, current_user).get("role") or "").strip().lower()


def active_user_required(
    view_function: Callable[P, ResponseReturnValue]
) -> Callable[P, ResponseReturnValue]:
    """
    Require a signed-in user whose account is still active in the database.

    Checking status on each request revokes access for accounts deactivated
    after login, including routes that do not otherwise need role checks.
    """

    @wraps(view_function)
    @login_required
    def decorated_view(*args: P.args, **kwargs: P.kwargs) -> ResponseReturnValue:
        user_id = get_current_user_id()
        if not user_id:
            return redirect(url_for("main.login"))

        try:
            current_user = fetch_one(
                """
                SELECT status
                FROM users
                WHERE user_id = %s
                LIMIT 1
                """,
                (user_id,)
            )
        except Exception as error:
            print(f"Active-account check failed: {error}")
            abort(503)

        account_status = str(
            (cast(Any, current_user) or {}).get("status") or ""
        ).strip().lower()

        if current_user is None or account_status != "active":
            logout_user()
            flash(
                "Your account is inactive or no longer available. "
                "Please contact the administrator if you need access.",
                "danger"
            )
            return redirect(url_for("main.login"))

        return view_function(*args, **kwargs)

    return decorated_view


def admin_required(
    view_function: Callable[P, ResponseReturnValue]
) -> Callable[P, ResponseReturnValue]:
    """
    Require an authenticated user whose current database role is Admin.

    The role is checked against the users table on every request instead of
    trusting a role value stored in the session. This means a role change in
    the database takes effect without waiting for the user's session to end.
    """

    @wraps(view_function)
    @active_user_required
    def decorated_view(*args: P.args, **kwargs: P.kwargs) -> ResponseReturnValue:
        user_id = get_current_user_id()

        if not user_id:
            # Keep the behavior consistent with the application's login flow.
            return redirect(url_for("main.login"))

        try:
            current_user = fetch_one(
                """
                SELECT role, status
                FROM users
                WHERE user_id = %s
                LIMIT 1
                """,
                (user_id,)
            )
        except Exception as error:
            # Fail closed: database/authorization errors must not grant access.
            print(f"Admin authorization check failed: {error}")
            abort(503)

        if current_user is None:
            # A deleted or invalid account must not retain privileged access.
            logout_user()
            return redirect(url_for("main.login"))

        # Re-check account status on every Admin-only request. This revokes
        # access immediately when an administrator deactivates an account.
        account_status = str(
            cast(Any, current_user).get("status") or ""
        ).strip().lower()
        if account_status != "active":
            logout_user()
            flash(
                "Your account is inactive. Please contact the administrator.",
                "danger"
            )
            return redirect(url_for("main.login"))

        current_role = str(
            cast(Any, current_user).get("role") or ""
        ).strip().lower()

        if current_role != "admin":
            abort(403)

        return view_function(*args, **kwargs)

    return decorated_view


# ============================================================
# PUBLIC / GENERAL ROUTES
# ============================================================

@main_bp.route("/")
def index():
    """
    InCraft default entry point.

    The root URL is the dashboard's canonical landing destination.
    Authenticated users are redirected directly to the dashboard.
    Visitors who are not signed in are sent to login and returned
    to the dashboard after successful authentication.
    """

    if get_current_user_id():
        return redirect(url_for("main.dashboard"))

    return redirect(
        url_for("main.login", next=url_for("main.dashboard"))
    )


# ============================================================
# AUTHENTICATION
# ============================================================

def _auth_database_available():
    """Return whether MySQL is reachable for authentication operations."""
    try:
        status = get_database_status()
        if status.get("connected"):
            return True

        # Keep detailed database information in the server log only.
        print(
            "Authentication database check failed: "
            f"{status.get('message', 'Unknown database connection error')}"
        )
    except Exception as error:
        print(f"Authentication database check failed: {error}")

    return False


@main_bp.route("/login", methods=["GET", "POST"])
def login():
    """
    User login with Admin / Staff role verification.
    """

    if request.method == "POST":

        # =========================================
        # GET FORM DATA
        # =========================================

        email = request.form.get(
            "email",
            ""
        ).strip().lower()

        password = request.form.get(
            "password",
            ""
        )

        selected_role = request.form.get(
            "role",
            ""
        ).strip().lower()


        # =========================================
        # VALIDATE REQUIRED FIELDS
        # =========================================

        if not email or not password or not selected_role:

            flash(
                "Please select your role and enter your email and password.",
                "danger"
            )

            return render_template(
                "login.html"
            )


        # =========================================
        # VALIDATE ROLE
        # =========================================

        # InCraft supports only Admin and Staff login roles.
        # A legacy Manager value in the database is not accepted for login.
        allowed_roles = {
            "admin",
            "staff"
        }

        if selected_role not in allowed_roles:

            flash(
                "Invalid account type selected.",
                "danger"
            )

            return render_template(
                "login.html"
            )


        # =========================================
        # CHECK DATABASE AVAILABILITY
        # =========================================

        if not _auth_database_available():
            flash(
                "The login service is temporarily unavailable. Please try again shortly.",
                "danger"
            )
            return render_template("login.html")

        # =========================================
        # FIND USER
        # =========================================

        user = fetch_one(
            """
            SELECT
                user_id,
                username,
                email,
                password,
                role,
                status
            FROM users
            WHERE LOWER(TRIM(email)) = %s
            LIMIT 1
            """,
            (email,)
        )


        # =========================================
        # USER NOT FOUND
        # =========================================

        if user is None:
            if not _auth_database_available():
                flash(
                    "The login service is temporarily unavailable. Please try again shortly.",
                    "danger"
                )
            else:
                # fetch_one() returns None for both no match and SQL errors.
                # db.py logs SQL errors to the server console.
                flash(
                    "Invalid email or password.",
                    "danger"
                )

            return render_template("login.html")


        # =========================================
        # CHECK ACCOUNT STATUS
        # =========================================

        account_status = str(
            cast(Any, user).get("status") or ""
        ).strip().lower()

        if account_status != "active":

            flash(
                "Your account is inactive. Please contact the administrator.",
                "danger"
            )

            return render_template(
                "login.html"
            )


        # =========================================
        # VERIFY SELECTED ROLE
        # =========================================

        actual_role = str(
            cast(Any, user).get(
                "role",
                ""
            )
        ).strip().lower()


        if actual_role != selected_role:

            if actual_role == "admin":

                flash(
                    "This account is registered as Admin. "
                    "Please select Admin to continue.",
                    "danger"
                )

            elif actual_role == "manager":

                flash(
                    "Manager accounts are no longer supported. "
                    "Please contact the administrator.",
                    "danger"
                )

            elif actual_role == "staff":

                flash(
                    "This account is registered as Staff. "
                    "Please select Staff to continue.",
                    "danger"
                )

            else:

                flash(
                    "Your account has an invalid role. "
                    "Please contact the administrator.",
                    "danger"
                )

            return render_template(
                "login.html"
            )


        # =========================================
        # GET STORED PASSWORD
        # =========================================

        stored_password = cast(Any, user).get(
            "password",
            ""
        )


        if not stored_password:

            flash(
                "Invalid email or password.",
                "danger"
            )

            return render_template(
                "login.html"
            )


        # =========================================
        # VERIFY PASSWORD
        # =========================================

        try:

            password_valid = check_password_hash(
                stored_password,
                password
            )

        except Exception:

            password_valid = False


        if not password_valid:

            flash(
                "Invalid email or password.",
                "danger"
            )

            return render_template(
                "login.html"
            )


        # =========================================
        # LOGIN USER
        # =========================================

        login_user(user)


        # =========================================
        # LOGIN ACTIVITY
        # =========================================

        # Activity logging is secondary to authentication. A logging failure
        # must not prevent a user with valid credentials from reaching the app.
        try:
            execute_query(
                """
                INSERT INTO activity_logs
                (
                    user_id,
                    action,
                    entity_type,
                    entity_id,
                    description,
                    created_at
                )
                VALUES (%s, %s, %s, %s, %s, %s)
                """,
                (
                    cast(Any, user).get("user_id"),
                    "LOGIN",
                    "USER",
                    cast(Any, user).get("user_id"),
                    f"{actual_role.capitalize()} user logged into InCraft",
                    datetime.now()
                )
            )
        except Exception as error:
            # Do not expose database details to the user.
            print(f"Login activity logging failed: {error}")


        # =========================================
        # SUCCESS MESSAGE
        # =========================================

        flash(
            f"Welcome back, {cast(Any, user).get('username')}!",
            "success"
        )


        # =========================================
        # REDIRECT
        # =========================================

        return redirect(
            url_for("main.dashboard")
        )


    # =========================================
    # GET REQUEST
    # =========================================

    return render_template(
        "login.html"
    )

@main_bp.route(
    "/register",
    methods=["GET", "POST"]
)
def register():
    """
    User registration.
    """

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        email = request.form.get(
            "email",
            ""
        ).strip().lower()

        password = request.form.get(
            "password",
            ""
        )

        confirm_password = request.form.get(
            "confirm_password",
            ""
        )

        if len(username) > 100:
            flash(
                "Username must be 100 characters or fewer.",
                "danger"
            )
            return render_template("register.html")

        if not username or not email or not password:

            flash(
                "All required fields must be completed.",
                "danger"
            )

            return render_template(
                "register.html"
            )

        if not validate_email(email):

            flash(
                "Please enter a valid email address.",
                "danger"
            )

            return render_template(
                "register.html"
            )

        if not validate_password(password):

            flash(
                "Password must contain at least 8 characters, "
                "including a letter and a number.",
                "danger"
            )

            return render_template(
                "register.html"
            )

        if password != confirm_password:

            flash(
                "Passwords do not match.",
                "danger"
            )

            return render_template(
                "register.html"
            )

        if not _auth_database_available():
            flash(
                "Registration is temporarily unavailable. Please try again shortly.",
                "danger"
            )
            return render_template("register.html")

        existing_user = fetch_one(
            """
            SELECT user_id
            FROM users
            WHERE LOWER(TRIM(email)) = %s
            LIMIT 1
            """,
            (email,)
        )

        if existing_user:

            flash(
                "An account with this email already exists.",
                "danger"
            )

            return render_template(
                "register.html"
            )

        password_hash = generate_password_hash(
            password
        )

        user_id = execute_query(
            """
            INSERT INTO users
            (
                username,
                email,
                password,
                role,
                status,
                created_at
            )
            VALUES (%s, %s, %s, %s, %s, %s)
            """,
            (
                username,
                email,
                password_hash,
                "staff",
                "active",
                datetime.now()
            )
        )

        if user_id is None:
            if not _auth_database_available():
                flash(
                    "Registration is temporarily unavailable because the database cannot be reached.",
                    "danger"
                )
            else:
                # execute_query() logs the underlying SQL error in db.py.
                # Common causes include schema/column mismatches and constraints.
                flash(
                    "Unable to create your account. Please check the server log and try again.",
                    "danger"
                )

            return render_template("register.html")

        flash(
            "Registration successful. You can now log in.",
            "success"
        )

        return redirect(
            url_for("main.login")
        )

    return render_template(
        "register.html"
    )


@main_bp.route("/logout", methods=["POST"])
def logout():
    """
    Log out the current user.
    """

    user_id = get_current_user_id()

    if user_id:
        # Logout must not fail because the secondary audit-log write failed.
        try:
            execute_query(
                """
                INSERT INTO activity_logs
                (
                    user_id,
                    action,
                    entity_type,
                    entity_id,
                    description,
                    created_at
                )
                VALUES (%s, %s, %s, %s, %s, %s)
                """,
                (
                    user_id,
                    "LOGOUT",
                    "USER",
                    user_id,
                    "User logged out of InCraft",
                    datetime.now()
                )
            )
        except Exception as error:
            # Keep internal database details out of the response.
            print(f"Logout activity logging failed: {error}")

    logout_user()

    flash(
        "You have been logged out.",
        "success"
    )

    return redirect(
        url_for("main.login")
    )


# ============================================================
# DASHBOARD
# ============================================================

@main_bp.route("/dashboard")
@active_user_required
def dashboard():
    """Render the dashboard with role-appropriate inventory and sales data."""

    current_user_id = get_current_user_id()
    current_role = get_current_user_role()
    if current_role is None:
        logout_user()
        return redirect(url_for("main.login"))

    is_admin = current_role == "admin"

    total_products_row = fetch_one(
        """
        SELECT COUNT(*) AS total
        FROM products
        WHERE status = 'active'
        """
    )

    total_stock_row = fetch_one(
        """
        SELECT COALESCE(SUM(quantity), 0) AS total
        FROM products
        WHERE status = 'active'
        """
    )

    if is_admin:
        total_sales_row = fetch_one(
            """
            SELECT COALESCE(SUM(total_amount), 0) AS total
            FROM sales
            """
        )
    else:
        total_sales_row = fetch_one(
            """
            SELECT COALESCE(SUM(total_amount), 0) AS total
            FROM sales
            WHERE user_id = %s
            """,
            (current_user_id,)
        )

    if is_admin:
        inventory_value_row = fetch_one(
            """
            SELECT COALESCE(SUM(quantity * price), 0) AS total
            FROM products
            WHERE status = 'active'
            """
        )
    else:
        # Inventory valuation is business-sensitive and is Admin-only.
        # None allows the template to distinguish "not permitted" from zero.
        inventory_value_row = None

    if is_admin:
        todays_sales_row = fetch_one(
            """
            SELECT COALESCE(SUM(total_amount), 0) AS total
            FROM sales
            WHERE DATE(sale_date) = CURDATE()
            """
        )
    else:
        todays_sales_row = fetch_one(
            """
            SELECT COALESCE(SUM(total_amount), 0) AS total
            FROM sales
            WHERE DATE(sale_date) = CURDATE()
              AND user_id = %s
            """,
            (current_user_id,)
        )

    low_stock_row = fetch_one(
        """
        SELECT COUNT(*) AS total
        FROM products
        WHERE status = 'active'
          AND quantity <= minimum_stock
        """
    )

    stock_distribution = fetch_one(
        """
        SELECT
            COALESCE(SUM(CASE WHEN quantity > minimum_stock THEN 1 ELSE 0 END), 0) AS stock_in_count,
            COALESCE(SUM(CASE WHEN quantity > 0 AND quantity <= minimum_stock THEN 1 ELSE 0 END), 0) AS stock_low_count,
            COALESCE(SUM(CASE WHEN quantity <= 0 THEN 1 ELSE 0 END), 0) AS stock_out_count
        FROM products
        WHERE status = 'active'
        """
    ) or {}

    low_stock_products = fetch_all(
        """
        SELECT
            name,
            product_code AS sku,
            quantity AS stock,
            minimum_stock AS reorder_level,
            price,
            category
        FROM products
        WHERE status = 'active'
          AND quantity <= minimum_stock
        ORDER BY quantity ASC, name ASC
        LIMIT 5
        """
    )

    recent_products = fetch_all(
        """
        SELECT
            name,
            product_code AS sku,
            category,
            quantity AS stock,
            price,
            minimum_stock AS reorder_level
        FROM products
        WHERE status = 'active'
        ORDER BY created_at DESC
        LIMIT 5
        """
    )

    if is_admin:
        recent_sales = fetch_all(
            """
            SELECT
                sale_id,
                invoice_number,
                customer_name,
                total_amount,
                sale_date
            FROM sales
            ORDER BY sale_date DESC
            LIMIT 5
            """
        )
    else:
        recent_sales = fetch_all(
            """
            SELECT
                sale_id,
                invoice_number,
                customer_name,
                total_amount,
                sale_date
            FROM sales
            WHERE user_id = %s
            ORDER BY sale_date DESC
            LIMIT 5
            """,
            (current_user_id,)
        )

    recent_activity_rows = fetch_all(
        """
        SELECT
            activity_id,
            action,
            entity_type,
            description,
            created_at
        FROM activity_logs
        WHERE user_id = %s
        ORDER BY created_at DESC
        LIMIT 8
        """,
        (get_current_user_id(),)
    )

    # Adapt activity-log records to the fields expected by dashboard.html.
    recent_activities = []
    for activity in recent_activity_rows or []:
        activity_row = cast(Any, activity)
        action = str(activity_row.get("action") or "Activity").replace("_", " ").strip()
        entity_type = str(activity_row.get("entity_type") or "").replace("_", " ").strip()
        description = activity_row.get("description") or ""
        created_at = activity_row.get("created_at")

        if action:
            title = action[:1].upper() + action[1:]
        else:
            title = "Activity"

        if not description and entity_type:
            description = entity_type[:1].upper() + entity_type[1:]

        if hasattr(created_at, "strftime"):
            activity_time = created_at.strftime("%d %b, %I:%M %p")
        else:
            activity_time = str(created_at or "")

        action_lower = action.lower()
        if "sale" in action_lower:
            icon = "🛒"
        elif "delete" in action_lower or "remove" in action_lower:
            icon = "🗑️"
        elif "update" in action_lower or "edit" in action_lower:
            icon = "✏️"
        elif "stock" in action_lower or "product" in action_lower or "add" in action_lower or "create" in action_lower:
            icon = "📦"
        else:
            icon = "📋"

        recent_activities.append({
            "icon": icon,
            "title": title,
            "description": description,
            "time": activity_time,
        })

    return render_template(
        "dashboard.html",
        total_products=(cast(Any, total_products_row) or {}).get("total", 0),
        total_stock=(cast(Any, total_stock_row) or {}).get("total", 0),
        total_sales=(cast(Any, total_sales_row) or {}).get("total", 0),
        inventory_value=(
            (cast(Any, inventory_value_row) or {}).get("total", 0)
            if is_admin
            else None
        ),
        is_admin=is_admin,
        todays_sales=(cast(Any, todays_sales_row) or {}).get("total", 0),
        low_stock=(cast(Any, low_stock_row) or {}).get("total", 0),
        stock_in_count=cast(Any, stock_distribution).get("stock_in_count", 0),
        stock_low_count=cast(Any, stock_distribution).get("stock_low_count", 0),
        stock_out_count=cast(Any, stock_distribution).get("stock_out_count", 0),
        low_stock_products=low_stock_products or [],
        recent_products=recent_products or [],
        recent_sales=recent_sales or [],
        recent_activity=recent_activity_rows or [],
        recent_activities=recent_activities,
    )


# ============================================================
# INVENTORY
# ============================================================

@main_bp.route("/inventory")
@active_user_required
def inventory():
    """
    Display inventory with search and stock information.
    """

    search = normalize_search_query(
        request.args.get(
            "search",
            ""
        )
    )

    if search:

        search_pattern = f"%{search}%"

        products = cast(list[dict[str, Any]], fetch_all(
            """
            SELECT
                product_id,
                product_code,
                name,
                category,
                quantity,
                price,
                minimum_stock,
                supplier,
                status,
                created_at
            FROM products
            WHERE status = 'active'
            AND (
                name LIKE %s
                OR product_code LIKE %s
                OR category LIKE %s
                OR supplier LIKE %s
            )
            ORDER BY created_at DESC
            """,
            (
                search_pattern,
                search_pattern,
                search_pattern,
                search_pattern
            )
        ))

    else:

        products = fetch_all(
            """
            SELECT
                product_id,
                product_code,
                name,
                category,
                quantity,
                price,
                minimum_stock,
                supplier,
                status,
                created_at
            FROM products
            WHERE status = 'active'
            ORDER BY created_at DESC
            """
        )

    # Database rows may be immutable RowType objects, so copy each row before
    # adding the computed fields used by the inventory template.
    products = [
        dict(cast(Any, product))
        for product in products
    ]

    for product in products:

        product["stock_status"] = get_stock_status(
            product.get(
                "quantity",
                0
            ),
            product.get(
                "minimum_stock",
                10
            )
        )

        product["stock_value"] = calculate_stock_value(
            product.get(
                "quantity",
                0
            ),
            product.get(
                "price",
                0
            )
        )

    return render_template(
        "inventory.html",
        products=products,
        search=search
    )


# ============================================================
# ADD PRODUCT
# ============================================================

@main_bp.route(
    "/inventory/add",
    methods=["GET", "POST"]
)
@admin_required
def add_product():
    """
    Add a new inventory product.
    """

    if request.method == "POST":

        product_code = request.form.get(
            "product_code",
            ""
        ).strip()

        name = request.form.get(
            "name",
            ""
        ).strip()

        category = request.form.get(
            "category",
            ""
        ).strip()

        quantity = request.form.get(
            "quantity",
            "0"
        )

        price = request.form.get(
            "price",
            "0"
        )

        minimum_stock = request.form.get(
            "minimum_stock",
            "10"
        )

        supplier = request.form.get(
            "supplier",
            ""
        ).strip()

        if (
            len(product_code) > 50
            or len(name) > 150
            or len(category) > 100
            or len(supplier) > 150
        ):
            flash(
                "Product fields exceed the allowed length.",
                "danger"
            )
            return render_template(
                "product_form.html",
                product=None
            )

        required, missing = validate_required_fields(
            {
                "product_code": product_code,
                "name": name,
                "category": category
            },
            [
                "product_code",
                "name",
                "category"
            ]
        )

        if not required:

            flash(
                "Please complete all required product fields.",
                "danger"
            )

            return render_template(
                "product_form.html",
                product=None
            )

        if (
            not str(quantity).strip().isdigit()
            or not validate_non_negative_number(quantity)
        ):

            flash(
                "Quantity must be a whole number zero or greater.",
                "danger"
            )

            return render_template(
                "product_form.html",
                product=None
            )

        if not validate_positive_number(
            price
        ):

            flash(
                "Price must be greater than zero.",
                "danger"
            )

            return render_template(
                "product_form.html",
                product=None
            )

        if (
            not str(minimum_stock).strip().isdigit()
            or not validate_non_negative_number(minimum_stock)
        ):

            flash(
                "Minimum stock must be a whole number zero or greater.",
                "danger"
            )

            return render_template(
                "product_form.html",
                product=None
            )

        existing_product = fetch_one(
            """
            SELECT product_id
            FROM products
            WHERE product_code = %s
            LIMIT 1
            """,
            (product_code,)
        )

        if existing_product:

            flash(
                "A product with this product code already exists.",
                "danger"
            )

            return render_template(
                "product_form.html",
                product=None
            )

        product_id = execute_query(
            """
            INSERT INTO products
            (
                product_code,
                name,
                category,
                quantity,
                price,
                minimum_stock,
                supplier,
                status,
                created_at,
                updated_at
            )
            VALUES
            (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            (
                product_code,
                name,
                category,
                safe_int(quantity),
                safe_float(price),
                safe_int(minimum_stock),
                supplier,
                "active",
                datetime.now(),
                datetime.now()
            )
        )

        if product_id is None:

            flash(
                "Unable to add the product.",
                "danger"
            )

            return render_template(
                "product_form.html",
                product=None
            )

        execute_query(
            """
            INSERT INTO activity_logs
            (
                user_id,
                action,
                entity_type,
                entity_id,
                description,
                created_at
            )
            VALUES (%s, %s, %s, %s, %s, %s)
            """,
            (
                get_current_user_id(),
                "CREATE",
                "PRODUCT",
                product_id,
                create_activity_description(
                    "Product added",
                    "Product",
                    name
                ),
                datetime.now()
            )
        )

        flash(
            "Product added successfully.",
            "success"
        )

        return redirect(
            url_for("main.inventory")
        )

    return render_template(
        "product_form.html",
        product=None
    )


# ============================================================
# EDIT PRODUCT
# ============================================================

@main_bp.route(
    "/inventory/edit/<int:product_id>",
    methods=["GET", "POST"]
)
@admin_required
def edit_product(product_id):
    """
    Edit an existing inventory product.
    """

    product = fetch_one(
        """
        SELECT *
        FROM products
        WHERE product_id = %s
        LIMIT 1
        """,
        (product_id,)
    )

    if product is None:

        flash(
            "Product not found.",
            "danger"
        )

        return redirect(
            url_for("main.inventory")
        )

    if request.method == "POST":

        name = request.form.get(
            "name",
            ""
        ).strip()

        category = request.form.get(
            "category",
            ""
        ).strip()

        quantity = request.form.get(
            "quantity",
            "0"
        )

        price = request.form.get(
            "price",
            "0"
        )

        minimum_stock = request.form.get(
            "minimum_stock",
            "10"
        )

        supplier = request.form.get(
            "supplier",
            ""
        ).strip()

        if (
            len(name) > 150
            or len(category) > 100
            or len(supplier) > 150
        ):
            flash(
                "Product fields exceed the allowed length.",
                "danger"
            )
            return render_template(
                "product_form.html",
                product=product
            )

        if not name or not category:

            flash(
                "Product name and category are required.",
                "danger"
            )

            return render_template(
                "product_form.html",
                product=product
            )

        if (
            not str(quantity).strip().isdigit()
            or not validate_non_negative_number(quantity)
        ):

            flash(
                "Quantity must be a whole number zero or greater.",
                "danger"
            )

            return render_template(
                "product_form.html",
                product=product
            )

        if not validate_positive_number(
            price
        ):

            flash(
                "Price must be greater than zero.",
                "danger"
            )

            return render_template(
                "product_form.html",
                product=product
            )

        if (
            not str(minimum_stock).strip().isdigit()
            or not validate_non_negative_number(minimum_stock)
        ):

            flash(
                "Minimum stock must be a whole number zero or greater.",
                "danger"
            )

            return render_template(
                "product_form.html",
                product=product
            )

        update_result = execute_query(
            """
            UPDATE products
            SET
                name = %s,
                category = %s,
                quantity = %s,
                price = %s,
                minimum_stock = %s,
                supplier = %s,
                updated_at = %s
            WHERE product_id = %s
            """,
            (
                name,
                category,
                safe_int(quantity),
                safe_float(price),
                safe_int(minimum_stock),
                supplier,
                datetime.now(),
                product_id
            )
        )

        if update_result is None:

            flash(
                "Unable to update the product.",
                "danger"
            )

            return render_template(
                "product_form.html",
                product=product
            )

        execute_query(
            """
            INSERT INTO activity_logs
            (
                user_id,
                action,
                entity_type,
                entity_id,
                description,
                created_at
            )
            VALUES (%s, %s, %s, %s, %s, %s)
            """,
            (
                get_current_user_id(),
                "UPDATE",
                "PRODUCT",
                product_id,
                create_activity_description(
                    "Product updated",
                    "Product",
                    name
                ),
                datetime.now()
            )
        )

        flash(
            "Product updated successfully.",
            "success"
        )

        return redirect(
            url_for("main.inventory")
        )

    return render_template(
        "product_form.html",
        product=product
    )


# ============================================================
# DELETE / ARCHIVE PRODUCT
# ============================================================

@main_bp.route(
    "/inventory/delete/<int:product_id>",
    methods=["POST"]
)
@admin_required
def delete_product(product_id):
    """
    Archive a product instead of permanently deleting it.

    This preserves historical sales and activity records.
    """

    product = fetch_one(
        """
        SELECT name
        FROM products
        WHERE product_id = %s
        LIMIT 1
        """,
        (product_id,)
    )

    if product is None:

        return jsonify({
            "success": False,
            "message": "Product not found."
        }), 404

    result = execute_query(
        """
        UPDATE products
        SET
            status = 'inactive',
            updated_at = %s
        WHERE product_id = %s
        """,
        (
            datetime.now(),
            product_id
        )
    )

    if result is None:

        return jsonify({
            "success": False,
            "message": "Unable to archive product."
        }), 500

    execute_query(
        """
        INSERT INTO activity_logs
        (
            user_id,
            action,
            entity_type,
            entity_id,
            description,
            created_at
        )
        VALUES (%s, %s, %s, %s, %s, %s)
        """,
        (
            get_current_user_id(),
            "DELETE",
            "PRODUCT",
            product_id,
            create_activity_description(
                "Product archived",
                "Product",
                cast(Any, product)["name"]
            ),
            datetime.now()
        )
    )

    return jsonify({
        "success": True,
        "message": "Product archived successfully."
    })


# ============================================================
# PRODUCT SEARCH API
# ============================================================

@main_bp.route(
    "/api/products/search"
)
@active_user_required
def search_products():
    """
    Search products asynchronously.
    """

    search = normalize_search_query(
        request.args.get(
            "q",
            ""
        )
    )

    if not search:

        return jsonify({
            "success": True,
            "products": []
        })

    search_pattern = f"%{search}%"

    products = fetch_all(
        """
        SELECT
            product_id,
            product_code,
            name,
            category,
            quantity,
            price,
            minimum_stock
        FROM products
        WHERE status = 'active'
        AND (
            name LIKE %s
            OR product_code LIKE %s
            OR category LIKE %s
        )
        ORDER BY name ASC
        LIMIT 20
        """,
        (
            search_pattern,
            search_pattern,
            search_pattern
        )
    )

    return jsonify({
        "success": True,
        "products": products
    })



# ============================================================
# GLOBAL SEARCH API
# ============================================================

@main_bp.route("/api/search", methods=["GET"])
@active_user_required
def global_search():
    """
    Search products, sales, and activity records from MySQL.

    Returns a unified list of results for the navbar search.
    """

    search = normalize_search_query(
        request.args.get("q", "")
    )

    if not search:
        return jsonify({
            "success": True,
            "results": []
        })

    search_pattern = f"%{search}%"
    results = []

    # --------------------------------------------------------
    # PRODUCTS
    # --------------------------------------------------------
    products = fetch_all(
        """
        SELECT
            product_id,
            product_code,
            name,
            category
        FROM products
        WHERE status = 'active'
          AND (
              name LIKE %s
              OR product_code LIKE %s
              OR category LIKE %s
              OR supplier LIKE %s
          )
        ORDER BY name ASC
        LIMIT 5
        """,
        (
            search_pattern,
            search_pattern,
            search_pattern,
            search_pattern
        )
    )

    for product in products:
        product_data = cast(Any, product)
        results.append({
            "type": "Product",
            "title": product_data.get("name") or "Unnamed product",
            "subtitle": (
                f"Code: {product_data.get('product_code') or 'N/A'}"
                f" · {product_data.get('category') or 'Uncategorized'}"
            ),
            "url": url_for(
                "main.inventory",
                search=search
            )
        })

    # --------------------------------------------------------
    # SALES
    # --------------------------------------------------------
    # Sales search is scoped by the current role. Staff must never be able
    # to discover another user's invoices or customer names through this API.
    current_user_id = get_current_user_id()
    current_role = get_current_user_role()
    if current_role is None:
        return jsonify({
            "success": False,
            "message": "Please sign in again."
        }), 401

    if current_role == "admin":
        sales_records = fetch_all(
            """
            SELECT
                sale_id,
                invoice_number,
                customer_name,
                sale_date
            FROM sales
            WHERE invoice_number LIKE %s
               OR customer_name LIKE %s
            ORDER BY sale_date DESC
            LIMIT 5
            """,
            (
                search_pattern,
                search_pattern
            )
        )
    else:
        sales_records = fetch_all(
            """
            SELECT
                sale_id,
                invoice_number,
                customer_name,
                sale_date
            FROM sales
            WHERE user_id = %s
              AND (
                  invoice_number LIKE %s
                  OR customer_name LIKE %s
              )
            ORDER BY sale_date DESC
            LIMIT 5
            """,
            (
                current_user_id,
                search_pattern,
                search_pattern
            )
        )

    for sale in sales_records:
        sale_data = cast(Any, sale)
        results.append({
            "type": "Sale",
            "title": (
                sale_data.get("invoice_number")
                or f"Sale #{sale_data.get('sale_id')}"
            ),
            "subtitle": (
                "Customer: "
                f"{sale_data.get('customer_name') or 'Walk-in customer'}"
            ),
            "url": url_for("main.sales")
        })

    # --------------------------------------------------------
    # ACTIVITY HISTORY
    # --------------------------------------------------------
    # Only administrators can search activity across all users.
    # Other signed-in users are limited to their own activity records.
    # current_user_id/current_role were verified above for sales search.
    if current_role == "admin":
        activities = fetch_all(
            """
            SELECT
                activity_id,
                action,
                entity_type,
                description,
                created_at
            FROM activity_logs
            WHERE (
                action LIKE %s
                OR entity_type LIKE %s
                OR description LIKE %s
            )
            ORDER BY created_at DESC
            LIMIT 5
            """,
            (search_pattern, search_pattern, search_pattern)
        )
    else:
        activities = fetch_all(
            """
            SELECT
                activity_id,
                action,
                entity_type,
                description,
                created_at
            FROM activity_logs
            WHERE user_id = %s
              AND (
                  action LIKE %s
                  OR entity_type LIKE %s
                  OR description LIKE %s
              )
            ORDER BY created_at DESC
            LIMIT 5
            """,
            (
                current_user_id,
                search_pattern,
                search_pattern,
                search_pattern
            )
        )

    for activity in activities:
        activity_record = cast(Any, activity)
        results.append({
            "type": "Activity",
            "title": (
                activity_record.get("description")
                or activity_record.get("action")
                or "Activity record"
            ),
            "subtitle": (
                f"{activity_record.get('entity_type') or 'Record'}"
                f" · {activity_record.get('action') or 'Activity'}"
            ),
            "url": url_for("main.activity")
        })

    return jsonify({
        "success": True,
        "results": results[:15]
    })


# ============================================================
# INVENTORY IMPORT
# ============================================================

@main_bp.route(
    "/inventory/import",
    methods=["POST"]
)
@admin_required
def import_inventory():
    """
    Import inventory data from a CSV or Excel (.xlsx) file.

    Required columns:
        product_code, name, category, quantity, price

    Optional column:
        minimum_stock (defaults to 10), supplier (defaults to blank)
    """

    uploaded_file = request.files.get("file")

    if uploaded_file is None or not uploaded_file.filename:
        flash("Please select a file to import.", "danger")
        return redirect(url_for("main.inventory"))

    if not allowed_file(uploaded_file.filename):
        flash("Unsupported file format.", "danger")
        return redirect(url_for("main.inventory"))

    extension = uploaded_file.filename.rsplit(".", 1)[-1].lower()
    if extension not in {"csv", "xlsx"}:
        flash("Please upload a CSV or Excel (.xlsx) inventory file.", "warning")
        return redirect(url_for("main.inventory"))

    max_import_bytes = 5 * 1024 * 1024
    workbook = None

    try:
        import csv
        import io
        import math

        raw_content = uploaded_file.stream.read(max_import_bytes + 1)

        if len(raw_content) > max_import_bytes:
            flash("Import file is too large. Maximum size is 5 MB.", "danger")
            return redirect(url_for("main.inventory"))

        if not raw_content:
            flash("The selected file is empty.", "warning")
            return redirect(url_for("main.inventory"))

        if extension == "csv":
            content = raw_content.decode("utf-8-sig")
            reader = csv.DictReader(io.StringIO(content))

            if not reader.fieldnames:
                flash("The CSV file must contain a header row.", "warning")
                return redirect(url_for("main.inventory"))

            headers = [
                str(header or "").strip().lower()
                for header in reader.fieldnames
            ]
            source_rows = (
                {
                    headers[index]: value
                    for index, value in enumerate(row.values())
                    if index < len(headers)
                }
                for row in reader
            )
        else:
            from openpyxl import load_workbook  # type: ignore[reportMissingModuleSource]

            workbook = load_workbook(
                io.BytesIO(raw_content),
                read_only=True,
                data_only=True
            )
            worksheet = workbook.active

            if worksheet is None:
                flash("The Excel file has no active worksheet.", "warning")
                return redirect(url_for("main.inventory"))

            sheet_rows = worksheet.iter_rows(values_only=True)
            header_row = next(sheet_rows, None)

            if not header_row:
                flash("The Excel file is empty or has no header row.", "warning")
                return redirect(url_for("main.inventory"))

            headers = [
                str(value or "").strip().lower()
                for value in header_row
            ]
            source_rows = (
                dict(zip(headers, values))
                for values in sheet_rows
            )

        required_headers = {"product_code", "name", "category", "quantity", "price"}
        missing_headers = sorted(required_headers - set(headers))

        if missing_headers:
            flash(
                "Import file is missing required column(s): "
                + ", ".join(missing_headers)
                + ".",
                "danger"
            )
            return redirect(url_for("main.inventory"))

        rows = []
        seen_product_codes = set()
        invalid_rows = 0

        for row in source_rows:
            # Normalize keys and values for CSV and Excel consistently.
            normalized_row = {
                str(key or "").strip().lower(): value
                for key, value in row.items()
                if key is not None
            }

            product_code = str(normalized_row.get("product_code") or "").strip()
            name = str(normalized_row.get("name") or "").strip()
            category = str(normalized_row.get("category") or "").strip()

            if not product_code and not name and not category:
                # Ignore fully blank spreadsheet rows.
                continue

            if not product_code or not name or not category:
                invalid_rows += 1
                continue

            raw_quantity = str(normalized_row.get("quantity", "") or "").strip()
            raw_price = str(normalized_row.get("price", "") or "").strip()
            raw_minimum_stock = str(
                normalized_row.get("minimum_stock", 10) or 10
            ).strip()

            try:
                quantity_value = float(raw_quantity)
                price = float(raw_price)
                minimum_value = float(raw_minimum_stock)

                if (
                    not math.isfinite(quantity_value)
                    or not math.isfinite(price)
                    or not math.isfinite(minimum_value)
                    or not quantity_value.is_integer()
                    or not minimum_value.is_integer()
                ):
                    invalid_rows += 1
                    continue

                quantity = int(quantity_value)
                minimum_stock = int(minimum_value)
            except (TypeError, ValueError, OverflowError):
                invalid_rows += 1
                continue

            if quantity < 0 or price <= 0 or minimum_stock < 0:
                invalid_rows += 1
                continue

            supplier = str(normalized_row.get("supplier") or "").strip()

            if (
                len(product_code) > 50
                or len(name) > 150
                or len(category) > 100
                or len(supplier) > 150
            ):
                invalid_rows += 1
                continue

            normalized_code = product_code.casefold()
            if normalized_code in seen_product_codes:
                invalid_rows += 1
                continue

            seen_product_codes.add(normalized_code)

            rows.append((
                product_code,
                name,
                category,
                quantity,
                price,
                minimum_stock,
                supplier,
                "active",
                datetime.now(),
                datetime.now()
            ))

        if not rows:
            flash(
                "No valid inventory records were found. Check the file headers "
                "and make sure quantities are whole numbers and prices are greater than zero.",
                "warning"
            )
            return redirect(url_for("main.inventory"))

        # The database helper executes the batch as one operation.
        # Existing product-code conflicts fail the batch rather than partially importing it.
        imported_count = execute_many(
            """
            INSERT INTO products
            (
                product_code,
                name,
                category,
                quantity,
                price,
                minimum_stock,
                supplier,
                status,
                created_at,
                updated_at
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s)
            """,
            rows
        )

        if imported_count is None:
            flash(
                "Inventory import failed. No records were imported. "
                "Check whether any product codes already exist.",
                "danger"
            )
            return redirect(url_for("main.inventory"))

        try:
            execute_query(
                """
                INSERT INTO activity_logs
                (
                    user_id,
                    action,
                    entity_type,
                    description,
                    created_at
                )
                VALUES (%s, %s, %s, %s, %s)
                """,
                (
                    get_current_user_id(),
                    "IMPORT",
                    "INVENTORY",
                    f"Imported {imported_count} inventory records",
                    datetime.now()
                )
            )
        except Exception as error:
            # Import success must not be misreported because audit logging failed.
            print(f"Inventory import activity logging failed: {error}")

        message = f"{imported_count} inventory records imported successfully."
        if invalid_rows:
            message += f" Skipped {invalid_rows} invalid or duplicate row(s)."
        flash(message, "success")

    except UnicodeDecodeError:
        flash("Import failed: CSV file must use UTF-8 encoding.", "danger")
    except Exception as error:
        # Keep internal database/file details out of the user-facing message.
        print(f"Inventory import failed: {error}")
        flash(
            "Inventory import failed. Check the file format, required columns, "
            "product codes, and database connection, then try again.",
            "danger"
        )
    finally:
        if workbook is not None:
            try:
                workbook.close()
            except Exception:
                pass

    return redirect(url_for("main.inventory"))


# ============================================================
# INVENTORY EXPORT
# ============================================================

@main_bp.route(
    "/inventory/export"
)
@admin_required
def export_inventory() -> ResponseReturnValue:
    """
    Export active inventory as CSV.
    """

    products = fetch_all(
        """
        SELECT
            product_code,
            name,
            category,
            quantity,
            price,
            minimum_stock,
            supplier,
            status
        FROM products
        WHERE status = 'active'
        ORDER BY name ASC
        """
    )

    if not products:

        flash(
            "There is no inventory data to export.",
            "warning"
        )

        return redirect(
            url_for("main.inventory")
        )

    execute_query(
        """
        INSERT INTO activity_logs
        (
            user_id,
            action,
            entity_type,
            description,
            created_at
        )
        VALUES (%s, %s, %s, %s, %s)
        """,
        (
            get_current_user_id(),
            "EXPORT",
            "INVENTORY",
            "Inventory exported as CSV",
            datetime.now()
        )
    )

    return cast(
        ResponseReturnValue,
        generate_csv_response(
            products,
            "incraft_inventory.csv"
        )
    )


# ============================================================
# SALES
# ============================================================

@main_bp.route("/sales")
@active_user_required
def sales():
    """
    Display sales records and active inventory products for the Record Sale form.
    Admins can see all sales; Staff can see only sales they recorded.
    """

    current_user_id = get_current_user_id()
    current_role = get_current_user_role()
    if current_role is None:
        logout_user()
        return redirect(url_for("main.login"))

    if current_role == "admin":
        sales_records = fetch_all(
            """
            SELECT
                sale_id,
                invoice_number,
                customer_name,
                total_amount,
                payment_method,
                sale_date
            FROM sales
            ORDER BY sale_date DESC
            """
        )
    else:
        sales_records = fetch_all(
            """
            SELECT
                sale_id,
                invoice_number,
                customer_name,
                total_amount,
                payment_method,
                sale_date
            FROM sales
            WHERE user_id = %s
            ORDER BY sale_date DESC
            """,
            (current_user_id,)
        )

    products = fetch_all(
        """
        SELECT
            product_id,
            product_code,
            name,
            category,
            quantity,
            price,
            minimum_stock,
            supplier,
            status
        FROM products
        WHERE status = 'active'
        ORDER BY name ASC
        """
    )

    return render_template(
        "sales.html",
        sales=sales_records,
        products=products
    )


# ============================================================
# CREATE SALE
# ============================================================

@main_bp.route(
    "/sales/create",
    methods=["POST"]
)
@active_user_required
def create_sale():
    """
    Record a new sale and update product stock.

    The sale header, sale item, and inventory update
    are executed inside one transaction.
    """

    # A UUID identifies one logical sale attempt. Retries must reuse the same
    # key so a network retry cannot create a second sale.
    raw_idempotency_key = str(
        request.form.get("idempotency_key", "")
    ).strip()

    try:
        idempotency_key = str(UUID(raw_idempotency_key))
    except (ValueError, TypeError, AttributeError):
        return jsonify({
            "success": False,
            "message": "The sale request key is missing or invalid. Reopen the form and try again."
        }), 400

    # Fast path for an already-completed request. A second lookup after a
    # failed transaction below handles the concurrent-request race.
    existing_sale = fetch_one(
        """
        SELECT invoice_number, total_amount
        FROM sales
        WHERE idempotency_key = %s
        LIMIT 1
        """,
        (idempotency_key,)
    )

    if existing_sale is not None:
        existing_row = cast(Any, existing_sale)
        existing_total = existing_row["total_amount"]
        return jsonify({
            "success": True,
            "message": "This sale was already recorded.",
            "invoice_number": existing_row["invoice_number"],
            "total_amount": existing_total,
            "formatted_total": format_currency(existing_total),
            "duplicate_request": True
        })

    # Do not silently coerce malformed form values such as "2.5",
    # "abc", or negative numbers into integers. Sales quantities must be
    # whole positive numbers supplied explicitly by the client.
    raw_product_id = str(
        request.form.get("product_id", "")
    ).strip()
    raw_quantity = str(
        request.form.get("quantity", "")
    ).strip()

    if (
        not raw_product_id.isdigit()
        or not raw_quantity.isdigit()
    ):
        return jsonify({
            "success": False,
            "message": "Product and quantity must be valid whole numbers."
        }), 400

    product_id = int(raw_product_id)
    quantity = int(raw_quantity)

    customer_name = request.form.get(
        "customer_name",
        ""
    ).strip()

    if len(customer_name) > 150:
        return jsonify({
            "success": False,
            "message": "Customer name must be 150 characters or fewer."
        }), 400

    payment_method = request.form.get(
        "payment_method",
        "cash"
    ).strip().lower()

    allowed_payment_methods = {
        "cash",
        "card",
        "upi",
        "bank_transfer",
        "other"
    }

    if payment_method not in allowed_payment_methods:

        return jsonify({
            "success": False,
            "message": "Invalid payment method."
        }), 400

    if product_id <= 0 or quantity <= 0:

        return jsonify({
            "success": False,
            "message": "Invalid product or quantity."
        }), 400

    product = fetch_one(
        """
        SELECT
            product_id,
            name,
            quantity,
            price
        FROM products
        WHERE product_id = %s
        AND status = 'active'
        LIMIT 1
        """,
        (product_id,)
    )

    if product is None:

        return jsonify({
            "success": False,
            "message": "Product not found."
        }), 404

    product_row = cast(Any, product)

    current_stock = safe_int(
        product_row["quantity"]
    )

    if quantity > current_stock:

        return jsonify({
            "success": False,
            "message": "Insufficient stock."
        }), 400

    # Keep the database DECIMAL value intact until the total is calculated.
    # Converting to float first can introduce avoidable currency rounding errors.
    unit_price = product_row["price"]

    total_amount = calculate_sale_total(
        quantity,
        unit_price
    )

    invoice_number = generate_reference(
        "INV"
    )

    sale_date = datetime.now()

    queries = [

        (
            """
            INSERT INTO sales
            (
                invoice_number,
                customer_name,
                total_amount,
                payment_method,
                sale_date,
                user_id,
                idempotency_key
            )
            VALUES (%s, %s, %s, %s, %s, %s, %s)
            """,
            (
                invoice_number,
                customer_name or None,
                total_amount,
                payment_method,
                sale_date,
                get_current_user_id(),
                idempotency_key
            ),
            1
        ),

        (
            """
            INSERT INTO sale_items
            (
                sale_id,
                product_id,
                quantity,
                unit_price
            )
            SELECT
                sale_id,
                %s,
                %s,
                %s
            FROM sales
            WHERE invoice_number = %s
            LIMIT 1
            """,
            (
                product_id,
                quantity,
                unit_price,
                invoice_number
            ),
            1
        ),

        (
            """
            UPDATE products
            SET
                quantity = quantity - %s,
                updated_at = %s
            WHERE product_id = %s
            AND status = 'active'
            AND quantity >= %s
            """,
            (
                quantity,
                datetime.now(),
                product_id,
                quantity
            ),
            1
        ),

        # Keep the audit record atomic with the sale. If any statement fails,
        # execute_transaction() rolls back the sale header, sale item, stock
        # deduction, and activity log together.
        (
            """
            INSERT INTO activity_logs
            (
                user_id,
                action,
                entity_type,
                description,
                created_at
            )
            VALUES (%s, %s, %s, %s, %s)
            """,
            (
                get_current_user_id(),
                "CREATE",
                "SALE",
                f"Sale {invoice_number} recorded for "
                f"{product_row['name']}",
                datetime.now()
            ),
            1
        )
    ]

    success = execute_transaction(
        queries
    )

    if not success:

        # If another concurrent request with this same key committed first,
        # the unique index rejects this transaction. Return the committed sale
        # instead of treating a successful retry as a new failure.
        existing_sale = fetch_one(
            """
            SELECT invoice_number, total_amount
            FROM sales
            WHERE idempotency_key = %s
            LIMIT 1
            """,
            (idempotency_key,)
        )

        if existing_sale is not None:
            existing_row = cast(Any, existing_sale)
            existing_total = existing_row["total_amount"]
            return jsonify({
                "success": True,
                "message": "This sale was already recorded.",
                "invoice_number": existing_row["invoice_number"],
                "total_amount": existing_total,
                "formatted_total": format_currency(existing_total),
                "duplicate_request": True
            })

        # The conditional stock UPDATE in the transaction is the final
        # concurrency guard. It prevents overselling even when another
        # request changes stock after the initial availability check.
        return jsonify({
            "success": False,
            "message": (
                "Unable to record the sale. Stock may have changed; "
                "refresh the page and try again."
            )
        }), 500

    return jsonify({
        "success": True,
        "message": "Sale recorded successfully.",
        "invoice_number": invoice_number,
        "total_amount": total_amount,
        "formatted_total": format_currency(
            total_amount
        )
    })


# ============================================================
# SALES EXPORT
# ============================================================

@main_bp.route(
    "/sales/export"
)
@admin_required
def export_sales():
    """
    Export sales records as CSV.
    """

    sales_records = fetch_all(
        """
        SELECT
            invoice_number,
            customer_name,
            total_amount,
            payment_method,
            sale_date
        FROM sales
        ORDER BY sale_date DESC
        """
    )

    if not sales_records:

        flash(
            "There are no sales records to export.",
            "warning"
        )

        return redirect(
            url_for("main.sales")
        )

    execute_query(
        """
        INSERT INTO activity_logs
        (
            user_id,
            action,
            entity_type,
            description,
            created_at
        )
        VALUES (%s, %s, %s, %s, %s)
        """,
        (
            get_current_user_id(),
            "EXPORT",
            "SALES",
            "Sales records exported as CSV",
            datetime.now()
        )
    )

    return cast(
        ResponseReturnValue,
        generate_csv_response(
            sales_records,
            "incraft_sales.csv"
        )
    )


# ============================================================
# DASHBOARD SALES TREND API
# ============================================================

@main_bp.route("/api/dashboard/sales")
@active_user_required
def dashboard_sales():
    """
    Return daily sales totals for the dashboard chart.
    The days query parameter is limited to 7, 30, or 90.
    """
    try:
        days = safe_int(request.args.get("days", 7), 7)
        if days not in (7, 30, 90):
            days = 7

        # Include today and the preceding (days - 1) calendar days.
        start_date = (datetime.now().date() - timedelta(days=days - 1))
        current_role = get_current_user_role()
        if current_role is None:
            logout_user()
            return jsonify({
                "success": False,
                "message": "Please sign in again."
            }), 401

        if current_role == "admin":
            sales_rows = fetch_all(
                """
                SELECT
                    DATE(sale_date) AS sales_day,
                    COALESCE(SUM(total_amount), 0) AS total
                FROM sales
                WHERE sale_date >= %s
                  AND sale_date < DATE_ADD(%s, INTERVAL 1 DAY)
                GROUP BY DATE(sale_date)
                ORDER BY sales_day ASC
                """,
                (start_date, datetime.now().date())
            ) or []
        else:
            sales_rows = fetch_all(
                """
                SELECT
                    DATE(sale_date) AS sales_day,
                    COALESCE(SUM(total_amount), 0) AS total
                FROM sales
                WHERE sale_date >= %s
                  AND sale_date < DATE_ADD(%s, INTERVAL 1 DAY)
                  AND user_id = %s
                GROUP BY DATE(sale_date)
                ORDER BY sales_day ASC
                """,
                (start_date, datetime.now().date(), get_current_user_id())
            ) or []

        totals_by_day = {}
        for row in sales_rows:
            row_data = cast(Any, row)
            sales_day = row_data.get("sales_day")
            if hasattr(sales_day, "strftime"):
                key = sales_day.strftime("%Y-%m-%d")
            else:
                key = str(sales_day or "")[:10]
            if key:
                totals_by_day[key] = float(row_data.get("total") or 0)

        labels = []
        values = []
        total_revenue = 0.0
        for offset in range(days):
            day = start_date + timedelta(days=offset)
            key = day.strftime("%Y-%m-%d")
            labels.append(day.strftime("%d %b"))
            value = round(totals_by_day.get(key, 0.0), 2)
            values.append(value)
            total_revenue += value

        return jsonify({
            "success": True,
            "days": days,
            "labels": labels,
            "values": values,
            "total": round(total_revenue, 2)
        })
    except Exception:
        # Keep internal database details out of the API response.
        return jsonify({
            "success": False,
            "message": "Unable to load sales trend data."
        }), 500


# ============================================================
# SMART DECISION SUPPORT API
# ============================================================

@main_bp.route(
    "/api/dashboard/insights"
)
@admin_required
def dashboard_insights():
    """
    Generate data-driven Smart Decision Support insights.

    The analysis uses:
        - Current inventory levels
        - Minimum stock levels
        - Product prices
        - Recent sales quantities
        - Recent sales revenue
        - Monthly revenue trends

    Returns:
        JSON response containing six insight categories.
    """

    try:
        products = fetch_all(
            """
            SELECT
                product_id,
                product_code,
                name,
                category,
                quantity,
                price,
                minimum_stock
            FROM products
            WHERE status = 'active'
            ORDER BY name ASC
            """
        )

        product_sales = fetch_all(
            """
            SELECT
                p.product_id,
                p.name,
                p.category,
                COALESCE(SUM(si.quantity), 0) AS quantity_sold,
                COALESCE(
                    SUM(si.quantity * si.unit_price),
                    0
                ) AS revenue
            FROM products p
            LEFT JOIN sale_items si
                ON p.product_id = si.product_id
            LEFT JOIN sales s
                ON si.sale_id = s.sale_id
            WHERE
                p.status = 'active'
                AND (
                    s.sale_id IS NULL
                    OR s.sale_date >= DATE_SUB(
                        NOW(),
                        INTERVAL 90 DAY
                    )
                )
            GROUP BY
                p.product_id,
                p.name,
                p.category
            ORDER BY quantity_sold DESC
            """
        )

        monthly_sales = fetch_all(
            """
            SELECT
                DATE_FORMAT(
                    sale_date,
                    '%Y-%m'
                ) AS month,
                COALESCE(
                    SUM(total_amount),
                    0
                ) AS revenue
            FROM sales
            WHERE sale_date >= DATE_SUB(
                NOW(),
                INTERVAL 6 MONTH
            )
            GROUP BY month
            ORDER BY month ASC
            """
        )

        sales_lookup = {}

        for row in product_sales:
            row = cast(dict[str, Any], row)
            sales_lookup[row["product_id"]] = row

        insights = [
            {
                "type": "restock",
                "priority": "low",
                "title": "Restock Recommendation",
                "message": (
                    "No products currently require "
                    "immediate restocking."
                ),
                "metric": "Stock levels healthy",
                "data": {}
            },
            {
                "type": "sales",
                "priority": "low",
                "title": "Sales Opportunity",
                "message": (
                    "No significant sales opportunity "
                    "was detected."
                ),
                "metric": "No strong signal",
                "data": {}
            },
            {
                "type": "slow-moving",
                "priority": "low",
                "title": "Slow-Moving Products",
                "message": (
                    "No slow-moving inventory "
                    "was detected."
                ),
                "metric": "Inventory turnover healthy",
                "data": {}
            },
            {
                "type": "revenue",
                "priority": "low",
                "title": "Revenue Insight",
                "message": (
                    "Not enough sales data to identify "
                    "a revenue trend."
                ),
                "metric": "Insufficient data",
                "data": {}
            },
            {
                "type": "forecast",
                "priority": "low",
                "title": "Demand Forecast",
                "message": (
                    "Not enough historical sales data "
                    "for a reliable demand indication."
                ),
                "metric": "Insufficient data",
                "data": {}
            },
            {
                "type": "optimization",
                "priority": "low",
                "title": "Inventory Optimization",
                "message": (
                    "Inventory appears reasonably balanced."
                ),
                "metric": "No major imbalance",
                "data": {}
            }
        ]

        # --------------------------------------------------------
        # 1. RESTOCK RECOMMENDATION
        # --------------------------------------------------------

        restock_products = []

        for product_row in products:
            product = cast(dict[str, Any], product_row)

            quantity = safe_int(
                product.get("quantity", 0)
            )

            minimum_stock = safe_int(
                product.get("minimum_stock", 0)
            )

            if quantity <= minimum_stock:
                shortage = max(
                    minimum_stock - quantity,
                    0
                )

                recommended_quantity = max(
                    minimum_stock * 2 - quantity,
                    1
                )

                restock_products.append({
                    "name": product.get(
                        "name",
                        "Unknown Product"
                    ),
                    "quantity": quantity,
                    "minimum_stock": minimum_stock,
                    "recommended_quantity": (
                        recommended_quantity
                    ),
                    "shortage": shortage
                })

        if restock_products:
            restock_products.sort(
                key=lambda item: (
                    item["shortage"],
                    item["minimum_stock"]
                ),
                reverse=True
            )

            item = restock_products[0]

            if item["quantity"] == 0:
                priority = "high"
            elif item["quantity"] < item["minimum_stock"]:
                priority = "high"
            else:
                priority = "medium"

            insights[0] = {
                "type": "restock",
                "priority": priority,
                "title": "Restock Recommendation",
                "message": (
                    f"{item['name']} is at "
                    f"{item['quantity']} units, "
                    f"which is at or below its "
                    f"minimum stock level of "
                    f"{item['minimum_stock']}. "
                    f"Consider adding approximately "
                    f"{item['recommended_quantity']} "
                    f"units."
                ),
                "metric": (
                    f"{len(restock_products)} "
                    f"product(s) need attention"
                ),
                "data": {
                    "products": restock_products
                }
            }

        # --------------------------------------------------------
        # 2. SALES OPPORTUNITY
        # --------------------------------------------------------

        sales_candidates = []

        for row in product_sales:
            row = cast(dict[str, Any], row)
            quantity_sold = safe_int(
                row["quantity_sold"]
            )

            revenue = safe_float(
                row["revenue"]
            )

            product_id = row["product_id"]

            product = next(
                (
                    cast(dict[str, Any], p)
                    for p in products
                    if cast(dict[str, Any], p).get(
                        "product_id"
                    ) == product_id
                ),
                None
            )

            if product is None:
                continue

            stock = safe_int(
                product.get("quantity", 0)
            )

            if quantity_sold > 0 and stock > 0:
                sales_candidates.append({
                    "name": row.get(
                        "name", "Unknown Product"
                    ) if row["name"] else "Unknown Product",
                    "quantity_sold": quantity_sold,
                    "revenue": revenue,
                    "stock": stock
                })

        if sales_candidates:
            sales_candidates.sort(
                key=lambda item: (
                    item["quantity_sold"],
                    item["revenue"]
                ),
                reverse=True
            )

            item = sales_candidates[0]

            insights[1] = {
                "type": "sales",
                "priority": "medium",
                "title": "Sales Opportunity",
                "message": (
                    f"{item['name']} has recorded "
                    f"{item['quantity_sold']} units "
                    f"sold during the recent analysis "
                    f"period. Current stock is "
                    f"{item['stock']} units."
                ),
                "metric": (
                    f"{item['quantity_sold']} units sold"
                ),
                "data": item
            }

        # --------------------------------------------------------
        # 3. SLOW-MOVING PRODUCTS
        # --------------------------------------------------------

        slow_products = []

        for product in products:
            product = cast(dict[str, Any], product)
            product_id = product.get("product_id")

            sales = sales_lookup.get(
                product_id,
                {}
            )

            quantity_sold = safe_int(
                sales.get("quantity_sold", 0)
            )

            stock = safe_int(
                product.get("quantity", 0)
            )

            if stock > 0 and quantity_sold <= 2:
                stock_value = round(
                    stock * safe_float(
                        product.get("price", 0)
                    ),
                    2
                )

                slow_products.append({
                    "name": product.get(
                        "name",
                        "Unknown Product"
                    ),
                    "stock": stock,
                    "quantity_sold": quantity_sold,
                    "stock_value": stock_value
                })

        if slow_products:
            slow_products.sort(
                key=lambda item: item["stock_value"],
                reverse=True
            )

            item = slow_products[0]

            insights[2] = {
                "type": "slow-moving",
                "priority": "medium",
                "title": "Slow-Moving Products",
                "message": (
                    f"{item['name']} currently has "
                    f"{item['stock']} units in stock "
                    f"but only {item['quantity_sold']} "
                    f"units were sold during the "
                    f"recent analysis period. "
                    f"Consider reviewing its stock level."
                ),
                "metric": (
                    f"{len(slow_products)} "
                    f"slow-moving product(s)"
                ),
                "data": {
                    "products": slow_products
                }
            }

        # --------------------------------------------------------
        # 4. REVENUE INSIGHT
        # --------------------------------------------------------

        revenue_rows = [
            cast(dict[str, Any], row)
            for row in product_sales
        ]

        total_revenue = sum(
            safe_float(
                row["revenue"]
            )
            for row in revenue_rows
        )

        revenue_products = [
            row
            for row in revenue_rows
            if safe_float(
                row["revenue"]
            ) > 0
        ]

        if revenue_products:
            revenue_products.sort(
                key=lambda row: safe_float(
                    row["revenue"]
                ),
                reverse=True
            )

            top_revenue = revenue_products[0]

            insights[3] = {
                "type": "revenue",
                "priority": "low",
                "title": "Revenue Insight",
                "message": (
                    f"{top_revenue['name'] or 'Product'} "
                    f"generated the highest product-level "
                    f"revenue during the recent analysis "
                    f"period."
                ),
                "metric": (
                    f"₹{safe_float(top_revenue['revenue']):,.2f}"
                ),
                "data": {
                    "product": top_revenue["name"] or "Unknown Product",
                    "revenue": safe_float(
                        top_revenue["revenue"]
                    ),
                    "total_revenue": total_revenue
                }
            }

        # --------------------------------------------------------
        # 5. DEMAND FORECAST
        # --------------------------------------------------------

        monthly_values = [
            safe_float(
                cast(dict[str, Any], row)["revenue"]
            )
            for row in monthly_sales
        ]

        if len(monthly_values) >= 2:
            recent_revenue = monthly_values[-1]
            previous_revenue = monthly_values[-2]

            if previous_revenue > 0:
                change_percent = (
                    (
                        recent_revenue -
                        previous_revenue
                    )
                    / previous_revenue
                ) * 100

                if change_percent > 10:
                    forecast_message = (
                        "Recent revenue increased "
                        f"by {change_percent:.1f}% "
                        "compared with the previous "
                        "month. Demand may continue "
                        "to require closer stock monitoring."
                    )

                    forecast_priority = "medium"

                elif change_percent < -10:
                    forecast_message = (
                        "Recent revenue decreased "
                        f"by {abs(change_percent):.1f}% "
                        "compared with the previous "
                        "month. Review demand before "
                        "increasing inventory."
                    )

                    forecast_priority = "medium"

                else:
                    forecast_message = (
                        "Recent revenue has remained "
                        "relatively stable compared "
                        "with the previous month."
                    )

                    forecast_priority = "low"

                insights[4] = {
                    "type": "forecast",
                    "priority": forecast_priority,
                    "title": "Demand Forecast",
                    "message": forecast_message,
                    "metric": (
                        f"{change_percent:+.1f}% "
                        "month-over-month"
                    ),
                    "data": {
                        "current_revenue": (
                            recent_revenue
                        ),
                        "previous_revenue": (
                            previous_revenue
                        ),
                        "change_percent": round(
                            change_percent,
                            2
                        )
                    }
                }

        # --------------------------------------------------------
        # 6. INVENTORY OPTIMIZATION
        # --------------------------------------------------------

        total_inventory_value = 0
        excess_stock_value = 0
        excess_products = []

        for product_row in products:
            product = cast(dict[str, Any], product_row)

            quantity = safe_int(
                product.get("quantity", 0)
            )

            minimum_stock = safe_int(
                product.get("minimum_stock", 0)
            )

            price = safe_float(
                product.get("price", 0)
            )

            stock_value = quantity * price

            total_inventory_value += stock_value

            sales = sales_lookup.get(
                product.get("product_id"),
                {}
            )

            quantity_sold = safe_int(
                sales.get("quantity_sold", 0)
            )

            if (
                minimum_stock > 0
                and quantity >= minimum_stock * 3
                and quantity_sold <= 2
            ):
                excess_products.append({
                    "name": product.get(
                        "name",
                        "Unknown Product"
                    ),
                    "stock": quantity,
                    "minimum_stock": minimum_stock,
                    "stock_value": round(
                        stock_value,
                        2
                    )
                })

                excess_stock_value += stock_value

        if excess_products:
            excess_products.sort(
                key=lambda item: item["stock_value"],
                reverse=True
            )

            insights[5] = {
                "type": "optimization",
                "priority": "medium",
                "title": "Inventory Optimization",
                "message": (
                    f"{len(excess_products)} product(s) "
                    "have relatively high stock compared "
                    "with their recent sales activity. "
                    "Review these products before "
                    "purchasing additional stock."
                ),
                "metric": (
                    f"₹{excess_stock_value:,.2f} "
                    "potentially tied in excess stock"
                ),
                "data": {
                    "products": excess_products,
                    "inventory_value": round(
                        total_inventory_value,
                        2
                    ),
                    "excess_stock_value": round(
                        excess_stock_value,
                        2
                    )
                }
            }

        return jsonify({
            "success": True,
            "insights": insights,
            "generated_at": datetime.now().isoformat()
        })

    except Exception as error:
        print(
            f"Smart insights error: {error}"
        )

        return jsonify({
            "success": False,
            "message": (
                "Unable to generate smart "
                "inventory insights."
            ),
            "insights": []
        }), 500


# ============================================================
# ANALYTICS
# ============================================================

@main_bp.route("/analytics")
@admin_required
def analytics():
    """
    Analytics page.
    """

    total_revenue = fetch_one(
        """
        SELECT
            COALESCE(
                SUM(total_amount),
                0
            ) AS total
        FROM sales
        """
    )

    total_transactions = fetch_one(
        """
        SELECT COUNT(*) AS total
        FROM sales
        """
    )

    inventory_value = fetch_one(
        """
        SELECT
            COALESCE(
                SUM(quantity * price),
                0
            ) AS total
        FROM products
        WHERE status = 'active'
        """
    )

    category_data = fetch_all(
        """
        SELECT
            category,
            COUNT(*) AS product_count,
            COALESCE(
                SUM(quantity),
                0
            ) AS total_stock,
            COALESCE(
                SUM(quantity * price),
                0
            ) AS stock_value
        FROM products
        WHERE status = 'active'
        GROUP BY category
        ORDER BY stock_value DESC
        """
    )

    monthly_sales = fetch_all(
        """
        SELECT
            DATE_FORMAT(
                sale_date,
                '%Y-%m'
            ) AS month,
            COUNT(*) AS transactions,
            COALESCE(
                SUM(total_amount),
                0
            ) AS revenue
        FROM sales
        GROUP BY month
        ORDER BY month ASC
        LIMIT 12
        """
    )

    top_products = fetch_all(
        """
        SELECT
            p.name,
            SUM(si.quantity) AS quantity_sold,
            SUM(
                si.quantity * si.unit_price
            ) AS revenue
        FROM sale_items si
        INNER JOIN products p
            ON p.product_id = si.product_id
        GROUP BY
            p.product_id,
            p.name
        ORDER BY quantity_sold DESC
        LIMIT 10
        """
    )

    low_stock_products = fetch_all(
        """
        SELECT
            product_id,
            product_code,
            name,
            quantity,
            minimum_stock
        FROM products
        WHERE status = 'active'
        AND quantity <= minimum_stock
        ORDER BY quantity ASC
        """
    )

    return render_template(
        "analytics.html",
        total_revenue=(
            cast(Any, total_revenue)["total"]
            if total_revenue
            else 0
        ),
        total_transactions=(
            cast(Any, total_transactions)["total"]
            if total_transactions
            else 0
        ),
        inventory_value=(
            cast(Any, inventory_value)["total"]
            if inventory_value
            else 0
        ),
        category_data=category_data,
        monthly_sales=monthly_sales,
        top_products=top_products,
        low_stock_products=low_stock_products
    )


# ============================================================
# ANALYTICS API
# ============================================================

@main_bp.route(
    "/api/analytics"
)
@admin_required
def analytics_api():
    """
    Return analytics data for JavaScript charts.
    """

    monthly_sales = fetch_all(
        """
        SELECT
            DATE_FORMAT(
                sale_date,
                '%Y-%m'
            ) AS month,
            COALESCE(
                SUM(total_amount),
                0
            ) AS revenue
        FROM sales
        GROUP BY month
        ORDER BY month ASC
        LIMIT 12
        """
    )

    category_data = fetch_all(
        """
        SELECT
            category,
            COALESCE(
                SUM(quantity * price),
                0
            ) AS value
        FROM products
        WHERE status = 'active'
        GROUP BY category
        ORDER BY value DESC
        """
    )

    return jsonify({
        "success": True,
        "monthly_sales": monthly_sales,
        "category_data": category_data
    })


# ============================================================
# ACTIVITY / HISTORY
# ============================================================

@main_bp.route("/activity")
@active_user_required
def activity():
    """
    Display user activity history.
    """

    # Activity history is private to the signed-in user by default.
    # Administrators may review the system-wide audit trail.
    current_user_id = get_current_user_id()
    current_user = fetch_one(
        "SELECT role FROM users WHERE user_id = %s LIMIT 1",
        (current_user_id,)
    )
    current_role = str((cast(Any, current_user) or {}).get("role") or "").strip().lower()

    if current_role == "admin":
        activities = fetch_all(
            """
            SELECT
                a.activity_id,
                a.action,
                a.entity_type,
                a.entity_id,
                a.description,
                a.created_at,
                u.username
            FROM activity_logs a
            LEFT JOIN users u
                ON u.user_id = a.user_id
            ORDER BY a.created_at DESC
            LIMIT 100
            """
        )
    else:
        activities = fetch_all(
            """
            SELECT
                a.activity_id,
                a.action,
                a.entity_type,
                a.entity_id,
                a.description,
                a.created_at,
                u.username
            FROM activity_logs a
            LEFT JOIN users u
                ON u.user_id = a.user_id
            WHERE a.user_id = %s
            ORDER BY a.created_at DESC
            LIMIT 100
            """,
            (current_user_id,)
        )

    return render_template(
        "activity.html",
        activities=activities
    )


# ============================================================
# SETTINGS
# ============================================================

@main_bp.route(
    "/settings",
    methods=["GET", "POST"]
)
@active_user_required
def settings():
    """
    User settings and password update.
    """

    user_id = get_current_user_id()

    user = fetch_one(
        """
        SELECT
            user_id,
            username,
            email,
            role
        FROM users
        WHERE user_id = %s
        LIMIT 1
        """,
        (user_id,)
    )

    if user is None:

        logout_user()

        return redirect(
            url_for("main.login")
        )

    if request.method == "POST":

        username = request.form.get(
            "username",
            ""
        ).strip()

        email = request.form.get(
            "email",
            ""
        ).strip().lower()

        current_password = request.form.get(
            "current_password",
            ""
        )

        new_password = request.form.get(
            "new_password",
            ""
        )

        confirm_password = request.form.get(
            "confirm_password",
            ""
        )

        if not username or not email:

            flash(
                "Username and email are required.",
                "danger"
            )

            return render_template(
                "settings.html",
                user=user
            )

        if not validate_email(email):

            flash(
                "Please enter a valid email address.",
                "danger"
            )

            return render_template(
                "settings.html",
                user=user
            )

        existing_email = fetch_one(
            """
            SELECT user_id
            FROM users
            WHERE LOWER(TRIM(email)) = %s
            AND user_id != %s
            LIMIT 1
            """,
            (
                email,
                user_id
            )
        )

        if existing_email:

            flash(
                "This email is already in use.",
                "danger"
            )

            return render_template(
                "settings.html",
                user=user
            )

        password_hash = None

        if new_password:

            if not current_password:

                flash(
                    "Current password is required.",
                    "danger"
                )

                return render_template(
                    "settings.html",
                    user=user
                )

            password_record = fetch_one(
                """
                SELECT password
                FROM users
                WHERE user_id = %s
                LIMIT 1
                """,
                (user_id,)
            )

            try:
                password_valid = (
                    password_record is not None
                    and bool(cast(Any, password_record)["password"])
                    and check_password_hash(
                        cast(Any, password_record)["password"],
                        current_password
                    )
                )
            except (TypeError, ValueError):
                password_valid = False

            if not password_valid:

                flash(
                    "Current password is incorrect.",
                    "danger"
                )

                return render_template(
                    "settings.html",
                    user=user
                )

            if not validate_password(
                new_password
            ):

                flash(
                    "New password must contain at least "
                    "8 characters, including a letter "
                    "and a number.",
                    "danger"
                )

                return render_template(
                    "settings.html",
                    user=user
                )

            if new_password != confirm_password:

                flash(
                    "New passwords do not match.",
                    "danger"
                )

                return render_template(
                    "settings.html",
                    user=user
                )

            password_hash = generate_password_hash(
                new_password
            )

        update_result = execute_query(
            """
            UPDATE users
            SET
                username = %s,
                email = %s,
                password = COALESCE(%s, password),
                updated_at = %s
            WHERE user_id = %s
            """,
            (
                username,
                email,
                password_hash,
                datetime.now(),
                user_id
            )
        )

        if update_result is None:

            flash(
                "Unable to update your settings.",
                "danger"
            )

            return render_template(
                "settings.html",
                user=user
            )

        session["username"] = username
        session["email"] = email

        execute_query(
            """
            INSERT INTO activity_logs
            (
                user_id,
                action,
                entity_type,
                entity_id,
                description,
                created_at
            )
            VALUES (%s, %s, %s, %s, %s, %s)
            """,
            (
                user_id,
                "UPDATE",
                "SETTINGS",
                user_id,
                "User settings updated",
                datetime.now()
            )
        )

        flash(
            "Settings updated successfully.",
            "success"
        )

        return redirect(
            url_for("main.settings")
        )

    return render_template(
        "settings.html",
        user=user
    )



# ============================================================
# ML-ASSISTED INVENTORY APIs (ADDITIVE; EXISTING ROUTES UNCHANGED)
# ============================================================

def _load_ml_inventory_data():
    """Load active products and up to 180 days of product-level unit sales."""
    products = fetch_all(
        """
        SELECT product_id, product_code, name, category, quantity,
               minimum_stock, supplier
        FROM products
        WHERE status = 'active'
        ORDER BY name ASC
        """
    )
    sales_rows = fetch_all(
        """
        SELECT
            si.product_id,
            DATE(s.sale_date) AS sale_day,
            COALESCE(SUM(si.quantity), 0) AS units_sold
        FROM sales AS s
        INNER JOIN sale_items AS si ON si.sale_id = s.sale_id
        WHERE s.sale_date >= DATE_SUB(CURDATE(), INTERVAL 180 DAY)
          AND s.sale_date < DATE_ADD(CURDATE(), INTERVAL 1 DAY)
        GROUP BY si.product_id, DATE(s.sale_date)
        ORDER BY sale_day ASC
        """
    )
    history_by_product = {}
    for row in sales_rows:
        item = dict(cast(Any, row))
        history_by_product.setdefault(item.get("product_id"), []).append({
            "sale_day": item.get("sale_day"),
            "units_sold": item.get("units_sold", 0),
        })
    return [dict(cast(Any, p)) for p in products], history_by_product


def _get_ml_lead_time_days():
    """Read a configurable default lead time without requiring a schema change."""
    import os
    try:
        return max(1, min(180, int(os.getenv("INCRAFT_DEFAULT_LEAD_TIME_DAYS", "7"))))
    except (TypeError, ValueError):
        return 7


def _build_ml_forecasts():
    products, history_by_product = _load_ml_inventory_data()
    return [
        forecast_product(product, history_by_product.get(product.get("product_id"), []))
        for product in products
    ]


@main_bp.route("/api/ml/forecast")
@admin_required
def ml_demand_forecast():
    """Return product-level unit-demand forecasts for the next 7 and 30 days."""
    try:
        forecasts = _build_ml_forecasts()
        return jsonify({
            "success": True,
            "forecast_horizon_days": [7, 30],
            "count": len(forecasts),
            "forecasts": forecasts,
        })
    except Exception as error:
        print(f"ML demand forecast failed: {error}")
        return jsonify({"success": False, "message": "Unable to generate demand forecasts."}), 500


@main_bp.route("/api/ml/stockout-risk")
@admin_required
def ml_stockout_risk():
    """Return estimated risk of stock running out before replenishment."""
    try:
        lead_time = _get_ml_lead_time_days()
        forecasts = _build_ml_forecasts()
        results = [assess_stockout(item, lead_time) for item in forecasts]
        results.sort(key=lambda item: ({"high": 0, "medium": 1, "low": 2}.get(item["risk_level"], 3), -item["risk_score"]))
        return jsonify({
            "success": True,
            "default_lead_time_days": lead_time,
            "lead_time_source": "INCRAFT_DEFAULT_LEAD_TIME_DAYS environment variable; defaults to 7 days",
            "count": len(results),
            "risks": results,
        })
    except Exception as error:
        print(f"ML stockout-risk analysis failed: {error}")
        return jsonify({"success": False, "message": "Unable to estimate stockout risk."}), 500


@main_bp.route("/api/ml/reorder-recommendations")
@admin_required
def ml_reorder_recommendations():
    """Return forecast-driven reorder recommendations; does not modify stock."""
    try:
        lead_time = _get_ml_lead_time_days()
        forecasts = _build_ml_forecasts()
        recommendations = [recommend_reorder(item, lead_time) for item in forecasts]
        recommendations.sort(key=lambda item: (not item["reorder_now"], {"urgent": 0, "high": 1, "medium": 2, "monitor": 3}.get(item["urgency"], 4)))
        return jsonify({
            "success": True,
            "default_lead_time_days": lead_time,
            "lead_time_source": "INCRAFT_DEFAULT_LEAD_TIME_DAYS environment variable; defaults to 7 days",
            "safety_stock_basis": "products.minimum_stock is used as a safety-stock proxy because no separate safety-stock field exists",
            "count": len(recommendations),
            "recommendations": recommendations,
        })
    except Exception as error:
        print(f"ML reorder recommendations failed: {error}")
        return jsonify({"success": False, "message": "Unable to generate reorder recommendations."}), 500


# ============================================================
# PRODUCT DETAILS API
# ============================================================

@main_bp.route(
    "/api/products/<int:product_id>"
)
@active_user_required
def product_details(product_id):
    """
    Return product information as JSON.
    """

    product = fetch_one(
        """
        SELECT
            product_id,
            product_code,
            name,
            category,
            quantity,
            price,
            minimum_stock,
            supplier,
            status,
            created_at,
            updated_at
        FROM products
        WHERE product_id = %s
          AND status = 'active'
        LIMIT 1
        """,
        (product_id,)
    )

    if product is None:

        return jsonify({
            "success": False,
            "message": "Product not found."
        }), 404

    product = dict(cast(Any, product))

    product["stock_status"] = get_stock_status(
        product["quantity"],
        product["minimum_stock"]
    )

    product["stock_value"] = calculate_stock_value(
        product["quantity"],
        product["price"]
    )

    return jsonify({
        "success": True,
        "product": product
    })


# ============================================================
# LOW STOCK API
# ============================================================

@main_bp.route(
    "/api/inventory/low-stock"
)
@active_user_required
def low_stock_api():
    """
    Return products requiring restocking.
    """

    products = fetch_all(
        """
        SELECT
            product_id,
            product_code,
            name,
            quantity,
            minimum_stock
        FROM products
        WHERE status = 'active'
        AND quantity <= minimum_stock
        ORDER BY quantity ASC
        """
    )

    return jsonify({
        "success": True,
        "count": len(products),
        "products": products
    })


# ============================================================
# HEALTH CHECK
# ============================================================

@main_bp.route(
    "/api/health"
)
def health_check():
    """
    Public deployment health check.

    Reports missing environment-variable names only. Never return
    database credentials, connection strings, or secret values.
    """

    missing_environment_variables = []

    if IS_PRODUCTION and not SECRET_KEY_CONFIGURED:
        missing_environment_variables.append("INCRAFT_SECRET_KEY")

    if IS_PRODUCTION and not DATABASE_CONFIGURED:
        missing_environment_variables.extend(MISSING_DATABASE_ENV_VARS)

    if missing_environment_variables:
        return jsonify({
            "application": "InCraft",
            "status": "configuration_required",
            "database": "not_configured",
            "missing_environment_variables": list(
                dict.fromkeys(missing_environment_variables)
            ),
            "message": (
                "Configure the listed environment variables in Vercel, "
                "then redeploy the application."
            ),
        }), 503

    # This endpoint is public. Do not expose raw database error messages,
    # connection details, or other internal diagnostics in its response.
    try:
        database_status = get_database_status()
        database_connected = bool(
            isinstance(database_status, dict)
            and database_status.get("connected")
        )
    except Exception as error:
        print(f"Health-check database probe failed: {error}")
        database_connected = False

    return jsonify({
        "application": "InCraft",
        "status": "running" if database_connected else "database_unavailable",
        "database": "connected" if database_connected else "unavailable",
        "secret_key_configured": SECRET_KEY_CONFIGURED,
    }), (200 if database_connected else 503)